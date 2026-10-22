"""Generate one Word probe for size-dependent OpenType MATH constructions."""

from pathlib import Path

from docx import Document
from docx.oxml import parse_xml
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
DESTINATION = ROOT / "fixtures/docx/probes/MathSizeVariants.docx"
MATH = "http://schemas.openxmlformats.org/officeDocument/2006/math"
WORD = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def math_run(value: str, size: int) -> str:
    return f'<m:r><w:rPr><w:sz w:val="{size * 2}"/></w:rPr><m:t>{value}</m:t></m:r>'


def add_case(document: Document, label: str, expression: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.add_run(f"{label}: Before ")
    paragraph._p.append(parse_xml(
        f'<m:oMath xmlns:m="{MATH}" xmlns:w="{WORD}">{expression}</m:oMath>'
    ))
    paragraph.add_run(" after")
    following = document.add_paragraph(f"{label}: Following paragraph")
    following.paragraph_format.space_after = Pt(2)


document = Document()
section = document.sections[0]
section.top_margin = section.bottom_margin = Inches(0.7)
section.left_margin = section.right_margin = Inches(0.85)
normal = document.styles["Normal"]
normal.font.name = "Aptos"
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(0)

document.add_heading("OpenType MATH size variants", 1)
document.add_paragraph(
    "Each expression is inline. Compare symbol size, spacing around the expression, "
    "and the baseline of the following paragraph."
)

for size in (10, 16, 20, 24):
    add_case(document, f"Radical {size} pt",
             f"<m:rad><m:deg/><m:e>{math_run('y', size)}</m:e></m:rad>")

for size in (10, 16, 20):
    add_case(document, f"Parentheses {size} pt",
             f"<m:d><m:dPr><m:begChr m:val=\"(\"/><m:endChr m:val=\")\"/>"
             f"</m:dPr><m:e>{math_run('y', size)}</m:e></m:d>")

for size in (10, 16, 20):
    add_case(document, f"Integral {size} pt",
             f"<m:nary><m:naryPr><m:chr m:val=\"&#x222B;\"/>"
             f"</m:naryPr><m:sub>{math_run('0', 10)}</m:sub>"
             f"<m:sup>{math_run('1', 10)}</m:sup>"
             f"<m:e>{math_run('y', size)}</m:e></m:nary>")

for length in (1, 6, 12):
    add_case(document, f"Overbar {length} letters",
             f"<m:acc><m:accPr><m:chr m:val=\"&#x0305;\"/>"
             f"</m:accPr><m:e>{math_run('x' * length, 12)}</m:e></m:acc>")

DESTINATION.parent.mkdir(parents=True, exist_ok=True)
document.save(DESTINATION)
print(DESTINATION)
