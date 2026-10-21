"""Generate a small Word probe for repeated spaces in styled paragraphs."""

from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt


root = Path(__file__).resolve().parents[2]
destination = root / "fixtures/docx/features/StyledRepeatedSpaces.docx"

document = Document()
section = document.sections[0]
section.top_margin = Inches(1)
section.left_margin = Inches(1)

style = document.styles["Subtitle"]
style.font.name = "Arial"
style.font.size = Pt(12)
style.paragraph_format.space_after = Pt(0)
style.paragraph_format.line_spacing = 1

for line in ("Alpha · Beta", "Alpha  · Beta", "Alpha ·  Beta"):
    document.add_paragraph(line, style="Subtitle")

table = document.add_table(rows=2, cols=1)
table.autofit = False
table.columns[0].width = Inches(1.15)
for cell, line in zip((table.cell(0, 0), table.cell(1, 0)),
                      ("AAAA BBBB Delta", "AAAA BBBB  Delta")):
    cell.width = Inches(1.15)
    paragraph = cell.paragraphs[0]
    paragraph.style = "Subtitle"
    paragraph.add_run(line)

for spaces in (" ", "  "):
    paragraph = document.add_paragraph(style="Subtitle")
    paragraph.paragraph_format.right_indent = Inches(5.7)
    paragraph.add_run("AAAA BBBB ")
    paragraph.add_run("CCCC" + spaces + "Omega")

destination.parent.mkdir(parents=True, exist_ok=True)
document.save(destination)
print(destination)
